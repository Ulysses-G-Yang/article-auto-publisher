"""Content Studio cover strategy, migration, and frozen-version tests."""

import asyncio
import io
import json
import sqlite3
import uuid
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from account_sessions.account_service import AccountSessionService
from account_sessions.database import AccountDatabase
from account_sessions.models import PlatformAccount
from account_sessions.permissions import LOCAL_WEB_CONTEXT
from content_studio.assets import AssetStore
from content_studio.contracts import (
    CoverInput,
    CreateDraftRequest,
    PatchDraftRequest,
    ReplaceTargetsRequest,
)
from content_studio.database import ContentDatabase, sqlite_url
from content_studio.errors import ContentAssetError, DraftValidationError
from content_studio.importers import DocxImportAdapter
from content_studio.models import ContentDraft, ContentVersion
from content_studio.service import ContentStudioService


def run(coroutine):
    return asyncio.run(coroutine)


def image_bytes(color: str = "red") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (24, 16), color).save(buffer, format="PNG")
    return buffer.getvalue()


def docx_bytes(*, with_image: bool) -> bytes:
    document = Document()
    document.core_properties.title = "导入封面标题"
    document.add_paragraph("正文段落")
    if with_image:
        document.add_picture(io.BytesIO(image_bytes()), width=Inches(0.3))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_service(tmp_path: Path) -> ContentStudioService:
    store = AssetStore(tmp_path / "assets")
    return ContentStudioService(
        ContentDatabase(sqlite_url(tmp_path / "content.db")),
        asset_store=store,
        docx_importer=DocxImportAdapter(store, work_root=tmp_path / "work"),
    )


def test_cover_input_is_strict() -> None:
    assert CoverInput(strategy="NONE").asset_id is None
    with pytest.raises(ValueError):
        CoverInput(strategy="EXPLICIT")
    with pytest.raises(ValueError):
        CoverInput(strategy="NONE", asset_id="a" * 36)


def test_old_content_schema_migrates_cover_columns_idempotently(tmp_path: Path) -> None:
    database_path = tmp_path / "legacy-content.db"
    with sqlite3.connect(database_path) as connection:
        connection.executescript(
            """
            CREATE TABLE content_drafts (
                draft_id TEXT PRIMARY KEY,
                seed_key TEXT UNIQUE,
                source_type TEXT NOT NULL,
                source_ref TEXT,
                title TEXT NOT NULL,
                blocks_json JSON NOT NULL,
                status TEXT NOT NULL,
                revision INTEGER NOT NULL,
                created_at DATETIME NOT NULL,
                updated_at DATETIME NOT NULL
            );
            CREATE TABLE content_versions (
                version_id TEXT PRIMARY KEY,
                draft_id TEXT NOT NULL,
                source_revision INTEGER NOT NULL,
                content_hash TEXT NOT NULL,
                title TEXT NOT NULL,
                blocks_json JSON NOT NULL,
                created_at DATETIME NOT NULL
            );
            INSERT INTO content_drafts VALUES
                ('draft-old', NULL, 'BLANK', NULL, '旧草稿', '[]', 'ACTIVE', 1,
                 '2026-08-01 00:00:00', '2026-08-01 00:00:00');
            INSERT INTO content_versions VALUES
                ('version-old', 'draft-old', 1, 'hash', '旧草稿', '[]',
                 '2026-08-01 00:00:00');
            """
        )

    database = ContentDatabase(sqlite_url(database_path))
    try:
        run(database.initialize())
        run(database.initialize())
        with sqlite3.connect(database_path) as connection:
            draft_columns = {
                row[1]: row[4]
                for row in connection.execute("PRAGMA table_info(content_drafts)")
            }
            version_columns = {
                row[1]: row[4]
                for row in connection.execute("PRAGMA table_info(content_versions)")
            }
        assert draft_columns["cover_strategy"] == "'NONE'"
        assert version_columns["cover_strategy"] == "'NONE'"
        assert "cover_asset_id" in draft_columns
        assert "cover_asset_id" in version_columns

        run(database.dispose())
        database = ContentDatabase(sqlite_url(database_path))
        run(database.initialize())

        async def read_rows():
            async with database.session() as session:
                return (
                    await session.get(ContentDraft, "draft-old"),
                    await session.get(ContentVersion, "version-old"),
                )

        draft, version = run(read_rows())
        assert draft.cover_strategy == "NONE"
        assert draft.cover_asset_id is None
        assert version.cover_strategy == "NONE"
        assert version.cover_asset_id is None
    finally:
        run(database.dispose())


def test_docx_import_defaults_cover_to_none_even_when_body_has_images(tmp_path: Path) -> None:
    service = make_service(tmp_path)

    async def scenario():
        await service.initialize()
        with_image = await service.import_docx(docx_bytes(with_image=True), "with-image.docx")
        without_image = await service.import_docx(
            docx_bytes(with_image=False), "without-image.docx"
        )
        await service.database.dispose()
        return with_image, without_image

    with_image, without_image = run(scenario())
    assert with_image["cover"] == {
        "strategy": "NONE",
        "asset_id": None,
        "asset_url": None,
    }
    assert without_image["cover"] == {
        "strategy": "NONE",
        "asset_id": None,
        "asset_url": None,
    }


def test_cover_patch_clear_reorder_and_fail_closed(tmp_path: Path) -> None:
    service = make_service(tmp_path)

    async def scenario():
        await service.initialize()
        draft = await service.create_draft(CreateDraftRequest(title="封面", blocks=[]))
        first = await service.add_asset(draft["draft_id"], image_bytes("red"), "first.png")
        second = await service.add_asset(draft["draft_id"], image_bytes("blue"), "second.png")
        first_blocks = [
            {"type": "image", "asset_id": first["asset_id"], "position": 0},
            {"type": "image", "asset_id": second["asset_id"], "position": 1},
        ]
        selected = await service.patch_draft(
            draft["draft_id"],
            PatchDraftRequest(
                revision=draft["revision"],
                title="封面",
                blocks=first_blocks,
                cover={"strategy": "FIRST_BODY_IMAGE"},
            ),
        )
        reordered = await service.patch_draft(
            draft["draft_id"],
            PatchDraftRequest(
                revision=selected["revision"],
                title="封面",
                blocks=[
                    {"type": "image", "asset_id": second["asset_id"], "position": 0},
                    {"type": "image", "asset_id": first["asset_id"], "position": 1},
                ],
            ),
        )
        with pytest.raises(DraftValidationError):
            await service.patch_draft(
                draft["draft_id"],
                PatchDraftRequest(
                    revision=reordered["revision"],
                    title="封面",
                    blocks=[{"type": "text", "text": "无图", "position": 0}],
                ),
            )
        cleared = await service.patch_draft(
            draft["draft_id"],
            PatchDraftRequest(
                revision=reordered["revision"],
                title="封面",
                blocks=[
                    {"type": "image", "asset_id": second["asset_id"], "position": 0},
                    {"type": "image", "asset_id": first["asset_id"], "position": 1},
                ],
                cover={"strategy": "NONE"},
            ),
        )
        await service.database.dispose()
        return selected, reordered, cleared

    selected, reordered, cleared = run(scenario())
    assert selected["cover"]["strategy"] == "FIRST_BODY_IMAGE"
    assert selected["cover"]["asset_id"] != reordered["cover"]["asset_id"]
    assert cleared["cover"] == {"strategy": "NONE", "asset_id": None, "asset_url": None}


def test_explicit_cover_only_asset_is_allowed_and_upload_does_not_select_it(
    tmp_path: Path,
) -> None:
    """封面可以是当前草稿的独立资产，不必重复出现在正文图文块中。"""

    service = make_service(tmp_path)

    async def scenario():
        await service.initialize()
        draft = await service.create_draft(
            CreateDraftRequest(
                title="独立封面",
                blocks=[{"type": "text", "text": "正文内容", "position": 0}],
            )
        )
        cover_asset = await service.add_asset(
            draft["draft_id"], image_bytes("purple"), "cover-only.png"
        )
        after_upload = await service.get_draft(draft["draft_id"])
        updated = await service.patch_draft(
            draft["draft_id"],
            PatchDraftRequest(
                revision=after_upload["revision"],
                title="独立封面",
                blocks=[{"type": "text", "text": "正文内容", "position": 0}],
                cover={"strategy": "EXPLICIT", "asset_id": cover_asset["asset_id"]},
            ),
        )
        await service.database.dispose()
        return after_upload, updated, cover_asset

    after_upload, updated, cover_asset = run(scenario())
    assert after_upload["cover"] == {
        "strategy": "NONE",
        "asset_id": None,
        "asset_url": None,
    }
    assert [block["type"] for block in updated["blocks"]] == ["text"]
    assert updated["cover"] == {
        "strategy": "EXPLICIT",
        "asset_id": cover_asset["asset_id"],
        "asset_url": f"/api/content-assets/{cover_asset['asset_id']}",
    }


def test_explicit_cover_must_belong_to_current_draft_and_upload_does_not_change_cover(
    tmp_path: Path,
) -> None:
    service = make_service(tmp_path)

    async def scenario():
        await service.initialize()
        first = await service.create_draft(CreateDraftRequest(title="一", blocks=[]))
        second = await service.create_draft(CreateDraftRequest(title="二", blocks=[]))
        asset = await service.add_asset(first["draft_id"], image_bytes(), "cover.png")
        before = await service.get_draft(second["draft_id"])
        with pytest.raises(ContentAssetError):
            await service.patch_draft(
                second["draft_id"],
                PatchDraftRequest(
                    revision=second["revision"],
                    title="二",
                    blocks=[],
                    cover={"strategy": "EXPLICIT", "asset_id": asset["asset_id"]},
                ),
            )
        after = await service.get_draft(second["draft_id"])
        own_asset = await service.add_asset(second["draft_id"], image_bytes("green"), "own.png")
        after_upload = await service.get_draft(second["draft_id"])
        assert after_upload["cover"] == before["cover"]
        updated = await service.patch_draft(
            second["draft_id"],
            PatchDraftRequest(
                revision=after["revision"],
                title="二",
                blocks=[],
                cover={"strategy": "EXPLICIT", "asset_id": own_asset["asset_id"]},
            ),
        )
        await service.database.dispose()
        return before, after, updated, own_asset

    before, after, updated, own_asset = run(scenario())
    for field in ("revision", "title", "blocks", "cover"):
        assert after[field] == before[field]
    assert updated["cover"] == {
        "strategy": "EXPLICIT",
        "asset_id": own_asset["asset_id"],
        "asset_url": f"/api/content-assets/{own_asset['asset_id']}",
    }
    assert "storage_path" not in json.dumps(updated, ensure_ascii=False)


def test_plan_freezes_cover_and_hash_changes_with_cover(tmp_path: Path) -> None:
    account_db = AccountDatabase(sqlite_url(tmp_path / "accounts.db"))
    account_id = str(uuid.uuid4())
    profile_root = tmp_path / "profiles"
    profile = profile_root / account_id
    profile.mkdir(parents=True)
    accounts = AccountSessionService(
        account_db,
        seed_legacy_profiles=False,
        allowed_profile_roots=(profile_root,),
    )
    service = ContentStudioService(
        ContentDatabase(sqlite_url(tmp_path / "content.db")),
        asset_store=AssetStore(tmp_path / "assets"),
        account_service=accounts,
    )

    async def scenario():
        await accounts.initialize()
        async with account_db.session() as session:
            session.add(
                PlatformAccount(
                    account_id=account_id,
                    platform="xiaoheihe",
                    platform_user_id="cover-user",
                    display_name="封面账号",
                    profile_path=str(profile),
                    status="ACTIVE",
                    session_status="VALID",
                    persist_login=True,
                )
            )
        await service.initialize()
        draft = await service.create_draft(CreateDraftRequest(title="冻结封面", blocks=[]))
        first = await service.add_asset(draft["draft_id"], image_bytes("red"), "one.png")
        second = await service.add_asset(draft["draft_id"], image_bytes("blue"), "two.png")
        selected = await service.patch_draft(
            draft["draft_id"],
            PatchDraftRequest(
                revision=draft["revision"],
                title="冻结封面",
                blocks=[
                    {"type": "image", "asset_id": first["asset_id"], "position": 0},
                    {"type": "image", "asset_id": second["asset_id"], "position": 1},
                ],
                cover={"strategy": "FIRST_BODY_IMAGE"},
            ),
        )
        targeted = await service.replace_targets(
            draft["draft_id"],
            ReplaceTargetsRequest(
                revision=selected["revision"],
                targets=[
                    {
                        "platform": "xiaoheihe",
                        "account_id": account_id,
                        "mode": "DRAFT",
                    }
                ],
            ),
            LOCAL_WEB_CONTEXT,
        )
        first_plan = await service.create_delivery_plan(
            draft["draft_id"], targeted["revision"], LOCAL_WEB_CONTEXT
        )
        first_context, _ = await service.get_plan_execution_context(
            first_plan["plan_id"], LOCAL_WEB_CONTEXT
        )
        first_resolved = await service.resolve_delivery_payload(first_context["version_id"])
        changed = await service.patch_draft(
            draft["draft_id"],
            PatchDraftRequest(
                revision=targeted["revision"],
                title="冻结封面",
                blocks=[
                    {"type": "image", "asset_id": first["asset_id"], "position": 0},
                    {"type": "image", "asset_id": second["asset_id"], "position": 1},
                ],
                cover={"strategy": "EXPLICIT", "asset_id": second["asset_id"]},
            ),
        )
        second_plan = await service.create_delivery_plan(
            draft["draft_id"], changed["revision"], LOCAL_WEB_CONTEXT
        )
        second_context, _ = await service.get_plan_execution_context(
            second_plan["plan_id"], LOCAL_WEB_CONTEXT
        )
        second_resolved = await service.resolve_delivery_payload(second_context["version_id"])
        first_public = await service.get_delivery_plan(
            first_plan["plan_id"], LOCAL_WEB_CONTEXT
        )
        await service.database.dispose()
        await account_db.dispose()
        return (
            first_plan,
            second_plan,
            first_public,
            first_context,
            first_resolved,
            second_resolved,
        )

    (
        first_plan,
        second_plan,
        first_public,
        first_context,
        first_resolved,
        second_resolved,
    ) = run(scenario())
    assert first_plan["cover"]["strategy"] == "FIRST_BODY_IMAGE"
    assert first_context["cover"] == first_plan["cover"]
    assert second_plan["cover"]["strategy"] == "EXPLICIT"
    assert first_plan["content_version"] != second_plan["content_version"]
    assert first_public["cover"] == first_plan["cover"]
    assert "storage_path" not in json.dumps(first_public, ensure_ascii=False)
    first_cover = first_resolved[3]
    second_cover = second_resolved[3]
    assert first_cover["strategy"] == "FIRST_BODY_IMAGE"
    assert first_cover["asset_id"] == first_plan["cover"]["asset_id"]
    assert first_cover["filename"] == "one.png"
    assert Path(str(first_cover["local_path"])).is_file()
    assert second_cover["strategy"] == "EXPLICIT"
    assert second_cover["asset_id"] == second_plan["cover"]["asset_id"]
    assert second_cover["filename"] == "two.png"
    assert Path(str(second_cover["local_path"])).is_file()

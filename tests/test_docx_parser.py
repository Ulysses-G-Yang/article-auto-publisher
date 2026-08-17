"""Generated DOCX regression tests for the parser/importer boundary."""

import asyncio
import io
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image

from content_studio.assets import AssetStore
from content_studio.importers import DocxImportAdapter
from core.docx_parser import DocxParser


def _save(document: Document, path: Path) -> Path:
    document.save(path)
    return path


def _image_path(tmp_path: Path) -> Path:
    path = tmp_path / "inline.png"
    Image.new("RGB", (16, 16), "red").save(path, format="PNG")
    return path


def test_core_title_precedes_later_heading_two(tmp_path: Path) -> None:
    document = Document()
    document.core_properties.title = "  Core Title\n"
    document.add_heading("Later heading", level=2)
    document.add_paragraph("Body text")
    path = _save(document, tmp_path / "core-title.docx")

    parser = DocxParser(images_dir=str(tmp_path / "images"))
    article = parser.parse(str(path))

    assert parser.get_title_from_article(article) == "Core Title"
    assert article.blocks[0].type == "heading"
    assert article.blocks[0].style_name == "Heading 2"


def test_title_style_and_no_core_title_fallback(tmp_path: Path) -> None:
    styled = Document()
    styled.core_properties.title = "   "
    styled.add_paragraph("Styled title", style="Title")
    styled.add_heading("Not the title", level=2)
    styled_path = _save(styled, tmp_path / "styled-title.docx")

    fallback = Document()
    fallback.add_heading("Not the body title", level=2)
    fallback.add_paragraph("  First body title  ")
    fallback_path = _save(fallback, tmp_path / "fallback-title.docx")

    parser = DocxParser(images_dir=str(tmp_path / "images"))
    assert parser.get_title_from_article(parser.parse(str(styled_path))) == "Styled title"
    assert parser.get_title_from_article(parser.parse(str(fallback_path))) == "First body title"


def test_adjacent_runs_keep_spaces_as_one_text_block(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("左侧 ")
    paragraph.add_run("有意义的")
    paragraph.add_run(" 空格")
    path = _save(document, tmp_path / "runs.docx")

    article = DocxParser(images_dir=str(tmp_path / "images")).parse(str(path))

    assert [(block.type, block.text) for block in article.blocks] == [
        ("text", "左侧 有意义的 空格")
    ]


def test_inline_image_splits_text_only_at_image_boundary(tmp_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("图片前 ")
    paragraph.add_run().add_picture(str(_image_path(tmp_path)), width=Inches(0.2))
    paragraph.add_run(" 图片后")
    path = _save(document, tmp_path / "inline-image.docx")

    article = DocxParser(images_dir=str(tmp_path / "images")).parse(str(path))

    assert [block.type for block in article.blocks] == ["text", "image", "text"]
    assert [block.text for block in article.blocks if block.type == "text"] == [
        "图片前 ",
        " 图片后",
    ]
    assert article.image_count == 1


def test_importer_removes_only_first_title_body_block(tmp_path: Path) -> None:
    document = Document()
    document.core_properties.title = "核心标题"
    document.add_paragraph("核心标题")
    document.add_heading("核心标题", level=2)
    document.add_paragraph("正文")
    source = io.BytesIO()
    document.save(source)

    adapter = DocxImportAdapter(
        AssetStore(tmp_path / "assets"),
        work_root=tmp_path / "work",
    )
    title, blocks, assets = asyncio.run(adapter.parse(source.getvalue(), "dedupe.docx"))

    assert title == "核心标题"
    assert assets == []
    assert [block["text"] for block in blocks] == ["核心标题", "正文"]
    assert [block["position"] for block in blocks] == [0, 1]

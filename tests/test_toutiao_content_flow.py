"""头条号 ContentVersion 图文顺序的离线端到端回归。

测试只在 pytest 临时目录生成 DOCX 和图片，不读取用户桌面文件、账号、
Cookie 或浏览器 Profile，也不会访问真实平台。
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import uuid
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

from docx import Document
from docx.shared import Inches
from PIL import Image

from content_studio.assets import AssetStore
from content_studio.content_document import (
    DELIVERY_POLICY_VERSION,
    normalize_for_delivery,
    project_to_v1,
)
from content_studio.database import ContentDatabase, sqlite_url
from content_studio.importers import DocxImportAdapter
from content_studio.models import ContentVersion
from content_studio.service import ContentStudioService
from platforms.toutiao import BODY_SELECTOR, TITLE_SELECTOR, ToutiaoPlatform

SOURCE_TITLE = "多设备共用显示器，别只盯KVM：先把输入源、USB-C和桌面空间理清楚"
EXPECTED_PLATFORM_TITLE = SOURCE_TITLE[:30]
EXPECTED_TOKENS = [
    ("P", "正文第01段"),
    ("P", "正文第02段"),
    ("I", ""),
    ("H2", "章节一"),
    ("P", "正文第03段"),
    ("P", "正文第04段"),
    ("I", ""),
    ("H2", "章节二"),
    ("P", "正文第05段"),
    ("I", ""),
    ("P", "正文第06段"),
    ("I", ""),
    ("H2", "章节三"),
    ("P", "正文第07段"),
    ("P", "正文第08段"),
    ("I", ""),
    ("H2", "章节四"),
    ("P", "正文第09段"),
    ("P", "正文第10段"),
    ("P", "正文第11段"),
    ("I", ""),
    ("P", "正文第12段"),
    ("I", ""),
    ("H2", "章节五"),
    ("P", "正文第13段"),
    ("P", "正文第14段"),
    ("P", "正文第15段"),
    ("P", "正文第16段"),
    ("P", "正文第17段"),
]


def _png_bytes(color: tuple[int, int, int]) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (36, 24), color=color).save(output, format="PNG")
    return output.getvalue()


def _toutiao_contract_docx() -> tuple[bytes, list[str]]:
    """构造与指定七图 Word 同形的 29-token 文档。"""

    image_payloads = [
        _png_bytes(color)
        for color in (
            (201, 41, 47),
            (25, 132, 72),
            (35, 90, 198),
            (218, 143, 31),
            (132, 63, 181),
            (35, 169, 176),
            (116, 84, 61),
        )
    ]
    document = Document()
    document.add_heading(SOURCE_TITLE, level=1)
    image_index = 0
    for kind, text in EXPECTED_TOKENS:
        if kind == "P":
            document.add_paragraph(text)
        elif kind == "H2":
            document.add_heading(text, level=2)
        else:
            document.add_paragraph().add_run().add_picture(
                io.BytesIO(image_payloads[image_index]),
                width=Inches(0.5),
            )
            image_index += 1
    output = io.BytesIO()
    document.save(output)
    return output.getvalue(), [
        hashlib.sha256(payload).hexdigest() for payload in image_payloads
    ]


class _NoLegacySource:
    def list_articles(self, *, limit: int = 50, offset: int = 0):
        return [], 0


class _FakeEditor:
    def __init__(self, page: _FakePage) -> None:
        self.page = page

    async def click(self, *, timeout: int) -> None:
        del timeout

    async def focus(self, *, timeout: int) -> None:
        del timeout

    async def count(self) -> int:
        return 1

    async def press(self, key: str) -> None:
        if key == "Backspace":
            self.page.tokens.clear()
            self.page.pending_block = True

    async def inner_text(self) -> str:
        return "\n".join(
            token["text"] for token in self.page.tokens if token["kind"] != "I"
        )

    def locator(self, selector: str):
        assert selector == ":scope > p:last-child"
        return _FakeTail(self.page)


class _FakeTail:
    def __init__(self, page: _FakePage) -> None:
        self.page = page

    async def count(self) -> int:
        return 1

    async def evaluate(self, script: str) -> bool:
        assert "breaks.length <= 1" in script
        assert "ProseMirror-trailingBreak" not in script
        return self.page.pending_block

    async def click(self, **_kwargs) -> None:
        return None


class _FakeTitleInput:
    def __init__(self, page: _FakePage) -> None:
        self.page = page

    async def click(self, *, timeout: int) -> None:
        del timeout

    async def count(self) -> int:
        return 1

    async def fill(self, value: str) -> None:
        self.page.title = value
        self.page.events.append("title")

    async def input_value(self) -> str:
        return self.page.title


class _FakeKeyboard:
    def __init__(self, page: _FakePage) -> None:
        self.page = page

    async def press(self, key: str) -> None:
        if key == "Enter":
            self.page.pending_block = True
        elif key == "Shift+Enter" and self.page.tokens:
            self.page.tokens[-1]["text"] += "\n"

    async def insert_text(self, value: str) -> None:
        if self.page.pending_block or not self.page.tokens:
            self.page.tokens.append({"kind": "P", "text": value})
            self.page.pending_block = False
        else:
            self.page.tokens[-1]["text"] += value


class _FakePage:
    def __init__(self) -> None:
        self.title = ""
        self.tokens: list[dict[str, str]] = []
        self.pending_block = True
        self.events: list[str] = []
        self.keyboard = _FakeKeyboard(self)

    def is_closed(self) -> bool:
        return False

    async def evaluate(self, script: str):
        if "const reuseExistingTail = true" in script:
            return self.pending_block
        if "requestAnimationFrame" in script:
            return None
        if "selection.rangeCount" in script:
            assert "syl-placeholder.ProseMirror-widget" in script
            assert "placeholders.length <= 1" in script
            return True
        raise AssertionError("未预期的页面脚本")

    def locator(self, selector: str):
        if selector == TITLE_SELECTOR:
            return _FakeTitleInput(self)
        if selector == BODY_SELECTOR:
            return _FakeEditor(self)
        raise AssertionError(f"未预期的 selector: {selector}")


class _CapturingToutiaoPlatform(ToutiaoPlatform):
    """保留生产 fill_title/fill_content，只隔离真实浏览器副作用。"""

    def __init__(self) -> None:
        super().__init__()
        self.page = _FakePage()
        self.simulator.random_delay = AsyncMock()
        self.uploaded_paths: list[Path] = []
        self.autosave_barrier_calls: list[dict[str, object]] = []
        self.restarted_draft_titles: list[str] = []
        self.initialization_snapshot_calls: list[dict[str, object]] = []
        self.write_window_waits = 0
        self.content_block_autosave_waits: list[int] = []

    async def _apply_h2_to_current_block(self) -> None:
        assert self.page.tokens
        self.page.tokens[-1]["kind"] = "H2"

    async def _upload_image(self, image_path: str) -> dict:
        assert self.page.title == EXPECTED_PLATFORM_TITLE
        path = Path(image_path)
        assert path.is_file()
        self.uploaded_paths.append(path)
        self.page.events.append("image")
        fingerprint = f"fixture.test/{hashlib.sha256(path.read_bytes()).hexdigest()}"
        self.page.tokens.append(
            {"kind": "I", "text": "", "fingerprint": fingerprint}
        )
        self.page.pending_block = False
        return {"success": True, "error": "", "fingerprint": fingerprint}

    async def _read_editor_tokens(self) -> list[dict[str, str]]:
        return deepcopy(self.page.tokens)

    async def _wait_for_autosave_barrier(self, **_kwargs) -> dict[str, object]:
        self.autosave_barrier_calls.append(dict(_kwargs))
        self._active_draft_id = "7665272216262623790"
        return {
            "generation": self._mutation_generation,
            "status": 200,
            "code": 0,
            "err_no": 0,
            "draft_ids": frozenset({self._active_draft_id}),
        }

    async def _restart_browser_for_initialized_draft(
        self,
        expected_title: str,
    ) -> None:
        self.restarted_draft_titles.append(expected_title)

    async def _wait_for_new_draft_write_window(self) -> None:
        self.write_window_waits += 1

    async def _wait_for_content_block_autosave(
        self,
        block_number: int,
    ) -> None:
        self.content_block_autosave_waits.append(block_number)

    async def _fetch_draft_snapshot(self, **kwargs):
        self.initialization_snapshot_calls.append(dict(kwargs))
        return SimpleNamespace(
            title_to_ids={
                EXPECTED_PLATFORM_TITLE: frozenset({self._active_draft_id})
            }
        )


def test_content_version_resolves_exact_toutiao_word_order(tmp_path: Path) -> None:
    payload, expected_image_hashes = _toutiao_contract_docx()
    asset_store = AssetStore(tmp_path / "assets")
    service = ContentStudioService(
        ContentDatabase(sqlite_url(tmp_path / "content.db")),
        asset_store=asset_store,
        legacy_source=_NoLegacySource(),
        docx_importer=DocxImportAdapter(
            asset_store,
            work_root=tmp_path / "work",
        ),
    )

    async def scenario():
        await service.initialize()
        imported = await service.import_docx(payload, "toutiao-contract.docx")
        delivery_document, loss_report = normalize_for_delivery(imported["document"])
        projection = project_to_v1(
            imported["document"],
            omit_title_block=True,
        ).blocks
        version_id = str(uuid.uuid4())
        async with service.database.session() as session:
            session.add(
                ContentVersion(
                    version_id=version_id,
                    draft_id=imported["draft_id"],
                    source_revision=imported["revision"],
                    content_hash="a" * 64,
                    title=imported["title"],
                    blocks_json=projection,
                    content_schema_version=2,
                    document_json=imported["document"],
                    delivery_document_json=delivery_document,
                    delivery_policy_version=DELIVERY_POLICY_VERSION,
                    delivery_loss_report_json=loss_report,
                    cover_strategy="NONE",
                    cover_asset_id=None,
                )
            )
        resolved = await service.resolve_delivery_payload(version_id)
        await service.database.dispose()
        return imported, resolved

    imported, (title, blocks, images, cover) = asyncio.run(scenario())
    adapter = _CapturingToutiaoPlatform()

    async def deliver_to_adapter():
        await adapter.fill_title(title)
        assert adapter.page.title == EXPECTED_PLATFORM_TITLE
        content_result = await adapter.fill_content(blocks, images)
        assert adapter.page.title == EXPECTED_PLATFORM_TITLE
        return content_result

    result = asyncio.run(deliver_to_adapter())

    image_hashes = iter(expected_image_hashes)
    expected_dict_tokens = []
    for kind, text in EXPECTED_TOKENS:
        token = {"kind": kind, "text": text}
        if kind == "I":
            token["fingerprint"] = f"fixture.test/{next(image_hashes)}"
        expected_dict_tokens.append(token)
    assert imported["title"] == SOURCE_TITLE
    assert title == SOURCE_TITLE
    assert adapter.page.title == EXPECTED_PLATFORM_TITLE
    assert len(adapter.page.title) == 30
    assert adapter.page.events == ["title"] + ["image"] * 7
    assert adapter.autosave_barrier_calls == [
        {
            "generation": 1,
            "stage": "标题初始化",
            "require_draft_id": True,
        }
    ]
    assert adapter.restarted_draft_titles == []
    assert adapter.write_window_waits == 0
    assert adapter.content_block_autosave_waits == list(range(1, 30))
    assert adapter.initialization_snapshot_calls == []
    assert len(blocks) == 29
    assert sum(block["type"] == "text" for block in blocks) == 17
    assert sum(block["type"] == "heading" for block in blocks) == 5
    assert sum(block["type"] == "image" for block in blocks) == 7
    assert [
        (
            "H2" if block["type"] == "heading" else "I" if block["type"] == "image" else "P",
            block.get("text", ""),
        )
        for block in blocks
    ] == EXPECTED_TOKENS
    assert adapter._expected_persisted_tokens == expected_dict_tokens
    assert adapter.page.tokens == expected_dict_tokens
    assert len(
        {
            token["fingerprint"]
            for token in adapter.page.tokens
            if token["kind"] == "I"
        }
    ) == 7
    assert result == {
        "text_ok": True,
        "expected_images": 7,
        "uploaded_images": 7,
        "failed_images": [],
        "media_status": "completed",
        "media_error": None,
    }
    assert cover["strategy"] == "NONE"
    assert cover["local_path"] is None
    assert len(images) == 7
    assert [
        hashlib.sha256(path.read_bytes()).hexdigest()
        for path in adapter.uploaded_paths
    ] == expected_image_hashes

""".docx 文档解析器 —— 提取文本 + 图片，保持顺序和位置关系"""

import json
import os
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from PIL import Image

from config import get_config

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
W_TEXT = f"{{{W_NS}}}t"
W_TAB = f"{{{W_NS}}}tab"
W_BREAK = f"{{{W_NS}}}br"
W_CARRIAGE_RETURN = f"{{{W_NS}}}cr"
A_BLIP = f"{{{A_NS}}}blip"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
WP_INLINE = f"{{{WP_NS}}}inline"
WP_ANCHOR = f"{{{WP_NS}}}anchor"
WP_DOC_PR = f"{{{WP_NS}}}docPr"
PIC_CNV_PR = "{http://schemas.openxmlformats.org/drawingml/2006/picture}cNvPr"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
M_OMATH = f"{{{M_NS}}}oMath"
M_OMATH_PARA = f"{{{M_NS}}}oMathPara"


@dataclass
class ContentBlock:
    """内容块：文本或图片"""

    type: str  # "text" | "image" | "heading"
    text: str | None = None
    style_name: str | None = None
    image_filename: str | None = None
    image_path: str | None = None
    image_width: int | None = None
    image_height: int | None = None
    position: int = 0


@dataclass
class ParsedArticle:
    """解析后的文章"""

    filename: str
    original_path: str
    blocks: list[ContentBlock] = field(default_factory=list)
    plain_text: str = ""
    image_count: int = 0
    char_count: int = 0
    document_title: str | None = None


@dataclass
class ParsedDocumentV2:
    """富结构解析的内部结果。

    ``document_v2`` 里的 ``__source_path`` 只在解析临时目录中存在，
    由 ``DocxImportAdapter.parse_v2`` 在校验前替换为受控 asset UUID。
    该 dataclass 不用于数据库、HTTP 或日志输出。
    """

    filename: str
    title: str
    v1_blocks: list[ContentBlock]
    document_v2: dict[str, Any]
    image_paths: list[Path] = field(default_factory=list)


class DocxParser:
    def __init__(self, images_dir: str | None = None):
        cfg = get_config()
        self.images_dir = images_dir or cfg["paths"]["images"]
        os.makedirs(self.images_dir, exist_ok=True)

    def parse(self, filepath: str) -> ParsedArticle:
        """解析 docx 文件，返回结构化文章"""
        filename = os.path.basename(filepath)
        doc = Document(filepath)

        article = ParsedArticle(
            filename=filename,
            original_path=filepath,
            document_title=self._normalize_title(doc.core_properties.title) or None,
        )

        # 创建文章专属图片目录
        article_images_dir = os.path.join(self.images_dir, os.path.splitext(filename)[0])
        os.makedirs(article_images_dir, exist_ok=True)

        # 遍历 document body 的所有子元素，保持原始顺序
        body = doc.element.body
        position = 0
        text_parts = []

        # 建立图片关系映射: rId -> image_part
        image_parts = {}
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                image_parts[rel.rId] = rel.target_part

        for child in body:
            # 处理段落
            if child.tag.endswith("}p"):
                para = self._find_paragraph(doc, child)
                if para is None:
                    continue

                # Walk the run XML in document order.  Text runs are accumulated
                # until an inline image boundary, so spaces at run boundaries are
                # preserved and adjacent runs become one text block.
                paragraph_parts = self._paragraph_parts(para, image_parts)

                # 部分 Word 生成器不会把纯文字暴露为标准 Run，保留段落级兜底。
                if not paragraph_parts and para.text.strip():
                    paragraph_parts.append(("text", para.text))

                style_name = para.style.name if para.style else None
                is_heading = self._is_heading_style(style_name)
                for part_type, value in paragraph_parts:
                    if part_type == "text":
                        block = ContentBlock(
                            type="heading" if is_heading else "text",
                            text=value,
                            style_name=para.style.name if para.style else None,
                            position=position,
                        )
                        article.blocks.append(block)
                        text_parts.append(value)
                        article.char_count += len(value)
                        position += 1
                        continue

                    image_part = value
                    img_filename = os.path.basename(image_part.partname)
                    img_path = os.path.join(article_images_dir, img_filename)
                    if not os.path.exists(img_path):
                        with open(img_path, "wb") as f:
                            f.write(image_part.blob)
                    width, height = self._get_image_size(img_path)
                    article.blocks.append(
                        ContentBlock(
                            type="image",
                            image_filename=img_filename,
                            image_path=img_path,
                            image_width=width,
                            image_height=height,
                            position=position,
                        )
                    )
                    article.image_count += 1
                    position += 1

            # 处理表格
            elif child.tag.endswith("}tbl"):
                table_text = self._extract_table_text(child)
                if table_text.strip():
                    block = ContentBlock(
                        type="text",
                        text=table_text.strip(),
                        position=position,
                    )
                    article.blocks.append(block)
                    text_parts.append(table_text.strip())
                    article.char_count += len(table_text.strip())
                    position += 1

        article.plain_text = "\n".join(text_parts)
        return article

    def parse_document_v2(self, filepath: str) -> ParsedDocumentV2:
        """解析 DOCX 为富结构 v2 草稿的内部表示。

        这是独立于历史 ``parse()`` 的 rich path。图片节点暂时携带
        ``__source_path``，仅供 Content Studio 导入适配器在临时目录内读取；
        调用方不得把该内部结果直接写入数据库或响应。
        """

        filename = os.path.basename(filepath)
        doc = Document(filepath)
        article_images_dir = Path(self.images_dir) / Path(filename).stem
        article_images_dir.mkdir(parents=True, exist_ok=True)

        # 只把可读取的内嵌图片关系加入映射。外部图片关系没有可控的
        # ``target_part``，不能把 URL 当作本地图片继续处理；遇到这种关系时
        # ``_rich_image`` 会 fail closed，而不是静默丢失正文图片。
        image_parts: dict[str, Any] = {}
        for rel in doc.part.rels.values():
            if "image" not in rel.reltype:
                continue
            try:
                target_part = rel.target_part
            except (AttributeError, KeyError, ValueError):
                continue
            if target_part is not None and hasattr(target_part, "blob"):
                image_parts[rel.rId] = target_part
        block_counter = [0]
        image_counter = [0]
        image_paths: list[Path] = []
        rich_blocks: list[dict[str, Any]] = []
        current_list_key: tuple[bool, int] | None = None
        body_font_size = self._body_font_size_points(doc)
        body_font_color = self._body_font_color(doc)
        nonempty_paragraphs = [
            paragraph for paragraph in doc.paragraphs if paragraph.text.strip()
        ]
        visual_title_paragraph = self._visual_title_paragraph(
            nonempty_paragraphs,
            body_font_size,
            body_font_color,
            doc,
        )

        for child in doc.element.body:
            if child.tag == qn("w:p"):
                paragraph = Paragraph(child, doc)
                explicit_heading_level = self._heading_level(
                    paragraph.style.name if paragraph.style else None
                )
                inferred_heading_level = None
                if explicit_heading_level is None and child is visual_title_paragraph:
                    inferred_heading_level = 1
                elif explicit_heading_level is None:
                    inferred_heading_level = self._visual_heading_level(
                        paragraph,
                        body_font_size,
                        doc,
                    )
                block = self._rich_paragraph(
                    paragraph,
                    doc,
                    image_parts,
                    article_images_dir,
                    image_counter,
                    image_paths,
                    block_counter,
                    inferred_heading_level=inferred_heading_level,
                )
                list_info = self._list_info(paragraph, doc)
                if list_info is not None:
                    ordered, level = list_info
                    list_key = (ordered, level)
                    if current_list_key == list_key and rich_blocks:
                        rich_blocks[-1]["items"].append({"blocks": [block]})
                    else:
                        list_block = {
                            "kind": "list",
                            "block_id": self._next_block_id(block_counter),
                            "ordered": ordered,
                            "level": level,
                            "items": [{"blocks": [block]}],
                        }
                        rich_blocks.append(list_block)
                        current_list_key = list_key
                else:
                    rich_blocks.append(block)
                    current_list_key = None
            elif child.tag == qn("w:tbl"):
                rich_blocks.append(
                    self._rich_table(
                        child,
                        doc,
                        image_parts,
                        article_images_dir,
                        image_counter,
                        image_paths,
                        block_counter,
                    )
                )
                current_list_key = None

        # Word 通常把图片说明保存为紧邻图片段落的 ``Caption`` 段落。
        # 只有“Caption + 紧邻的纯图片段落”这一无歧义关系才折叠到 ImageNode；
        # 其它 Caption 段落保留为普通段落，避免误吞正文。
        rich_blocks = self._attach_safe_captions(rich_blocks)

        # 无样式稿件的可见主标题才是投递标题的权威来源。Word 的
        # core_properties.title 经常保留旧标题、缩写标题或编辑前版本；若先
        # 使用该元数据，再靠文本完全相等寻找正文块，可见 H1 只要多一个括号
        # 说明就会失去 title_block_id，并被错误地作为正文 H1 送进能力门。
        #
        # 因此先按已经完成的视觉/显式标题识别结果，在开头三个非空文本块中
        # 寻找唯一 H1，并直接固化它的 block_id。只有没有可靠可见 H1 时，
        # 才回退到 Word 元数据和原有精确匹配逻辑。
        visible_title = self._leading_title_block(rich_blocks)
        if visible_title is not None:
            title = self._normalize_title(self._rich_block_text(visible_title))
            title_block_id = visible_title.get("block_id")
        else:
            core_title = self._normalize_title(doc.core_properties.title)
            title = core_title or self._title_from_rich_blocks(rich_blocks, filename)
            title_block_id = self._find_title_block_id(rich_blocks, title)

        # Word 模板和人工排版经常把正文小节也错误套成 Heading 1。标题块已经
        # 通过 ``title_block_id`` 唯一绑定后，其余 H1 不再有歧义：它们属于
        # 正文层级，统一按 H2 保存，避免把模板样式错误传播到平台能力门。
        self._normalize_body_heading_one(rich_blocks, title_block_id)
        document_v2: dict[str, Any] = {
            "schema_version": 2,
            "title": title,
            "blocks": rich_blocks,
            "source_fidelity": "NATIVE",
        }
        if title_block_id is not None:
            document_v2["title_block_id"] = title_block_id

        return ParsedDocumentV2(
            filename=filename,
            title=title,
            v1_blocks=self._legacy_blocks_from_rich(rich_blocks),
            document_v2=document_v2,
            image_paths=image_paths,
        )

    # The short alias is useful to callers that already name import paths parse_v2.
    parse_v2 = parse_document_v2

    @staticmethod
    def _next_block_id(counter: list[int]) -> str:
        counter[0] += 1
        return f"doc-block-{counter[0]:06d}"

    def _rich_paragraph(
        self,
        paragraph: Paragraph,
        doc,
        image_parts: dict[str, Any],
        image_dir: Path,
        image_counter: list[int],
        image_paths: list[Path],
        block_counter: list[int],
        *,
        hyperlink_target: str | None = None,
        inferred_heading_level: int | None = None,
    ) -> dict[str, Any]:
        if any(node.tag in {M_OMATH, M_OMATH_PARA} for node in paragraph._p.iter()):
            raise ValueError("DOCX 公式无法安全映射到平台正文")
        style_name = paragraph.style.name if paragraph.style else None
        explicit_level = self._heading_level(style_name)
        level = explicit_level if explicit_level is not None else inferred_heading_level
        block: dict[str, Any] = {
            "kind": "heading" if level is not None else "paragraph",
            "block_id": self._next_block_id(block_counter),
            "children": [],
        }
        if style_name:
            block["style_name"] = style_name
        if level is not None:
            block["level"] = level

        for child in paragraph._p:
            if child.tag == qn("w:r"):
                block["children"].extend(
                    self._rich_run_children(
                        child,
                        paragraph,
                        image_parts,
                        image_dir,
                        image_counter,
                        image_paths,
                        hyperlink_target,
                    )
                )
            elif child.tag == qn("w:hyperlink"):
                target, link_title = self._hyperlink_target(doc, child)
                for run_element in child.iter(qn("w:r")):
                    block["children"].extend(
                        self._rich_run_children(
                            run_element,
                            paragraph,
                            image_parts,
                            image_dir,
                            image_counter,
                            image_paths,
                            target,
                            link_title,
                        )
                    )
        if explicit_level is None and inferred_heading_level is not None:
            # 原始 Normal 只是识别前的 Word 来源样式；晋级为 canonical heading
            # 后必须移除，不能制造“Normal + H2”的自相矛盾结构。视觉 marks
            # 仍完整保留在原始文档，平台投递副本再统一归一化。
            block.pop("style_name", None)
        return block

    def _rich_run_children(
        self,
        run_element,
        parent,
        image_parts: dict[str, Any],
        image_dir: Path,
        image_counter: list[int],
        image_paths: list[Path],
        hyperlink_target: str | None = None,
        hyperlink_title: str | None = None,
    ) -> list[dict[str, Any]]:
        run = Run(run_element, parent)
        marks = self._run_marks(run)
        children: list[dict[str, Any]] = []
        for node in run_element.iter():
            if node.tag == W_TEXT:
                text = node.text or ""
                if text:
                    text_node: dict[str, Any] = {"kind": "text", "text": text}
                    if marks:
                        text_node["marks"] = list(marks)
                    if hyperlink_target is not None:
                        text_node["link"] = {"href": hyperlink_target}
                        if hyperlink_title:
                            text_node["link"]["title"] = hyperlink_title
                    children.append(text_node)
            elif node.tag == W_TAB:
                children.append({"kind": "text", "text": "\t"})
            elif node.tag in {W_BREAK, W_CARRIAGE_RETURN}:
                children.append({"kind": "text", "text": "\n"})
            elif node.tag == A_BLIP:
                image = self._rich_image(
                    node,
                    run_element,
                    image_parts,
                    image_dir,
                    image_counter,
                    image_paths,
                )
                if image is not None:
                    children.append(image)
        return children

    def _rich_image(
        self,
        blip,
        run_element,
        image_parts: dict[str, Any],
        image_dir: Path,
        image_counter: list[int],
        image_paths: list[Path],
    ) -> dict[str, Any] | None:
        rel_id = blip.get(f"{{{R_NS}}}embed") or blip.get(f"{{{R_NS}}}link")
        image_part = image_parts.get(rel_id)
        if image_part is None:
            raise ValueError("DOCX 图片不是可控的内嵌资源")
        image_counter[0] += 1
        suffix = Path(str(image_part.partname)).suffix.lower() or ".img"
        image_path = image_dir / f"image-{image_counter[0]:04d}{suffix}"
        image_path.write_bytes(image_part.blob)
        image_paths.append(image_path)
        width, height = self._get_image_size(str(image_path))
        image: dict[str, Any] = {
            "kind": "image",
            "__source_path": str(image_path),
            "__original_filename": image_path.name,
            "anchor": {"kind": self._image_anchor_kind(run_element)},
        }
        alt = self._image_alt(run_element)
        image["alt"] = alt or image_path.name
        if width > 0:
            image["width"] = width
        if height > 0:
            image["height"] = height
        return image

    def _attach_safe_captions(self, blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
        """在明确相邻时把 Caption 段落绑定到前一个图片节点。

        该方法递归处理列表和表格单元格，但不会跨容器寻找图片。未满足
        “Caption 样式 + 紧邻纯图片段落”条件的说明段落原样保留。
        """

        normalized: list[dict[str, Any]] = []
        index = 0
        while index < len(blocks):
            block = blocks[index]
            kind = block.get("kind")
            if kind == "list":
                block = dict(block)
                block["items"] = [
                    {
                        **item,
                        "blocks": self._attach_safe_captions(item.get("blocks", [])),
                    }
                    for item in block.get("items", [])
                ]
            elif kind == "table":
                block = dict(block)
                block["rows"] = [
                    {
                        **row,
                        "cells": [
                            {
                                **cell,
                                "blocks": self._attach_safe_captions(cell.get("blocks", [])),
                            }
                            for cell in row.get("cells", [])
                        ],
                    }
                    for row in block.get("rows", [])
                ]

            is_caption = (
                kind in {"paragraph", "heading"}
                and str(block.get("style_name") or "").strip().casefold() == "caption"
            )
            caption_text = self._rich_block_text(block).strip() if is_caption else ""
            previous = normalized[-1] if normalized else None
            image_node = self._pure_image_node(previous)
            if caption_text and image_node is not None:
                image_node = dict(image_node)
                image_node["caption"] = caption_text
                previous = dict(previous)
                previous["children"] = [image_node]
                normalized[-1] = previous
                index += 1
                continue

            normalized.append(block)
            index += 1
        return normalized

    @staticmethod
    def _pure_image_node(block: dict[str, Any] | None) -> dict[str, Any] | None:
        if not block or block.get("kind") not in {"paragraph", "heading"}:
            return None
        children = block.get("children")
        if not isinstance(children, list) or len(children) != 1:
            return None
        child = children[0]
        return child if child.get("kind") == "image" else None

    @staticmethod
    def _image_anchor_kind(run_element) -> str:
        if any(node.tag == WP_ANCHOR for node in run_element.iter()):
            return "floating"
        return "inline"

    @staticmethod
    def _image_alt(run_element) -> str | None:
        for node in run_element.iter():
            if node.tag in {WP_DOC_PR, PIC_CNV_PR}:
                for attribute in ("descr", "title"):
                    value = node.get(attribute)
                    if value and value.strip():
                        return value.strip()
        return None

    @staticmethod
    def _run_marks(run: Run) -> tuple[str, ...]:
        marks: list[str] = []
        if run.bold:
            marks.append("bold")
        if run.italic:
            marks.append("italic")
        if run.underline:
            marks.append("underline")
        if run.font.strike:
            marks.append("strike")
        style_name = run.style.name if run.style else ""
        if "code" in " ".join(str(style_name).split()).casefold():
            marks.append("code")
        return tuple(marks)

    @classmethod
    def _body_font_size_points(cls, doc) -> float | None:
        """推断正文基准字号；证据不足时返回 ``None`` 并拒绝视觉标题晋级。

        只采集非全粗体段落中的普通文字 run，避免标题本身抬高基准。字号可以
        来自 run、字符样式或段落样式，但不会凭文件名、段落位置或固定 11pt
        猜测正文格式。
        """

        sizes: Counter[float] = Counter()
        paragraphs = [paragraph for paragraph in doc.paragraphs if paragraph.text.strip()]
        for paragraph in paragraphs[1:]:
            if cls._heading_level(
                paragraph.style.name if paragraph.style else None
            ) is not None:
                continue
            text_runs = [
                run for run in cls._paragraph_runs(paragraph) if run.text.strip()
            ]
            if not text_runs or all(cls._effective_bold(run, paragraph) for run in text_runs):
                continue
            for run in text_runs:
                if cls._effective_bold(run, paragraph):
                    continue
                size = cls._effective_font_size_points(run, paragraph, doc)
                if size is not None:
                    weight = cls._visible_character_count(run.text)
                    if weight:
                        sizes[round(size, 2)] += weight
        if not sizes:
            return None
        max_weight = max(sizes.values())
        # 同权时选更小字号，避免偶发的大字号提示语抬高正文基准。
        return float(min(size for size, weight in sizes.items() if weight == max_weight))

    @classmethod
    def _body_font_color(cls, doc) -> str | None:
        colors: Counter[str] = Counter()
        paragraphs = [paragraph for paragraph in doc.paragraphs if paragraph.text.strip()]
        for paragraph in paragraphs[1:]:
            if cls._heading_level(
                paragraph.style.name if paragraph.style else None
            ) is not None:
                continue
            for run in cls._paragraph_runs(paragraph):
                weight = cls._visible_character_count(run.text)
                if not weight or cls._effective_bold(run, paragraph):
                    continue
                color = cls._effective_font_color(run, paragraph)
                if color is not None:
                    colors[color] += weight
        if not colors:
            return None
        max_weight = max(colors.values())
        return sorted(color for color, weight in colors.items() if weight == max_weight)[0]

    @classmethod
    def _visual_title_paragraph(
        cls,
        paragraphs: list[Paragraph],
        body_font_size: float | None,
        body_font_color: str | None,
        doc,
    ):
        """按“前三非空 + 字号差/粗体”选择唯一视觉主标题。"""

        if not paragraphs or body_font_size is None:
            return None
        candidates: list[tuple[tuple[float, ...], Paragraph]] = []
        for rank, paragraph in enumerate(paragraphs[:3]):
            style_name = paragraph.style.name if paragraph.style else None
            if cls._heading_level(style_name) is not None:
                continue
            text = paragraph.text.strip()
            if not text or len(text) > 40 or "\n" in text or "\r" in text:
                continue
            if any(node.tag == A_BLIP for node in paragraph._p.iter()):
                continue
            font_size = cls._paragraph_font_size_points(paragraph, doc)
            bold_ratio = cls._paragraph_bold_ratio(paragraph)
            size_signal = font_size is not None and font_size >= body_font_size + 4.0
            bold_signal = bold_ratio >= 0.70
            # “位于前三”是第一条证据；还必须至少命中字号或粗体证据。
            if not (size_signal or bold_signal):
                continue
            color = cls._paragraph_font_color(paragraph)
            color_signal = color is not None and color != body_font_color
            score = (
                float(1 + int(size_signal) + int(bold_signal)),
                float(int(color_signal)),
                float((font_size or body_font_size) - body_font_size),
                bold_ratio,
                float(-rank),
            )
            candidates.append((score, paragraph))
        if not candidates:
            return None
        return max(candidates, key=lambda item: item[0])[1]._p

    @classmethod
    def _visual_heading_level(
        cls,
        paragraph: Paragraph,
        body_font_size: float | None,
        doc=None,
    ) -> int | None:
        """按字号或 70% 粗体规则把普通样式短段落识别为语义 H2。"""

        if body_font_size is None:
            return None
        text = paragraph.text.strip()
        if not text or len(text) > 30 or "\n" in text or "\r" in text:
            return None
        if re.search(r"[。！？!?；;]", text[:-1]):
            return None
        if any(node.tag == A_BLIP for node in paragraph._p.iter()):
            return None
        if paragraph._p.find(qn("w:hyperlink")) is not None:
            return None
        style_name = " ".join(
            str(paragraph.style.name if paragraph.style else "").split()
        ).casefold()
        if style_name == "caption":
            return None

        text_runs = [run for run in cls._paragraph_runs(paragraph) if run.text.strip()]
        if not text_runs:
            return None
        heading_size = cls._paragraph_font_size_points(paragraph, doc)
        bold_ratio = cls._paragraph_bold_ratio(paragraph)
        size_signal = heading_size is not None and heading_size > body_font_size + 0.01
        if not size_signal and bold_ratio < 0.70:
            return None
        return 2

    @classmethod
    def _paragraph_bold_ratio(cls, paragraph: Paragraph) -> float:
        total = 0
        bold = 0
        for run in cls._paragraph_runs(paragraph):
            weight = cls._visible_character_count(run.text)
            if not weight:
                continue
            total += weight
            if cls._effective_bold(run, paragraph):
                bold += weight
        return bold / total if total else 0.0

    @classmethod
    def _paragraph_font_size_points(
        cls,
        paragraph: Paragraph,
        doc=None,
    ) -> float | None:
        sizes: Counter[float] = Counter()
        for run in cls._paragraph_runs(paragraph):
            weight = cls._visible_character_count(run.text)
            if not weight:
                continue
            size = cls._effective_font_size_points(run, paragraph, doc)
            if size is not None:
                sizes[round(size, 2)] += weight
        if not sizes:
            return None
        max_weight = max(sizes.values())
        return float(max(size for size, weight in sizes.items() if weight == max_weight))

    @classmethod
    def _paragraph_font_color(cls, paragraph: Paragraph) -> str | None:
        colors: Counter[str] = Counter()
        for run in cls._paragraph_runs(paragraph):
            weight = cls._visible_character_count(run.text)
            if not weight:
                continue
            color = cls._effective_font_color(run, paragraph)
            if color is not None:
                colors[color] += weight
        if not colors:
            return None
        max_weight = max(colors.values())
        return sorted(color for color, weight in colors.items() if weight == max_weight)[0]

    @staticmethod
    def _visible_character_count(value: str) -> int:
        return sum(1 for char in str(value or "") if not char.isspace())

    @staticmethod
    def _paragraph_runs(paragraph: Paragraph) -> list[Run]:
        """返回段落 XML 中的全部 run，包括 hyperlink 内部 run。"""

        return [Run(element, paragraph) for element in paragraph._p.iter(qn("w:r"))]

    @staticmethod
    def _effective_bold(run: Run, paragraph: Paragraph) -> bool:
        for value in (
            run.bold,
            getattr(getattr(run, "style", None), "font", None).bold
            if getattr(run, "style", None) is not None
            else None,
            getattr(getattr(paragraph, "style", None), "font", None).bold
            if getattr(paragraph, "style", None) is not None
            else None,
        ):
            if value is not None:
                return bool(value)
        return False

    @classmethod
    def _effective_font_size_points(
        cls,
        run: Run,
        paragraph: Paragraph,
        doc=None,
    ) -> float | None:
        candidates = [run.font.size]
        run_style = getattr(run, "style", None)
        paragraph_style = getattr(paragraph, "style", None)
        candidates.extend(cls._style_font_sizes(run_style))
        candidates.extend(cls._style_font_sizes(paragraph_style))
        if doc is not None:
            try:
                candidates.extend(cls._style_font_sizes(doc.styles["Normal"]))
            except (KeyError, TypeError):
                pass
        for size in candidates:
            if size is not None:
                return float(size.pt)
        return None

    @staticmethod
    def _style_font_sizes(style) -> list[Any]:
        """按当前样式到 ``base_style`` 的顺序返回显式字号。

        python-docx 不会把基于样式继承的字号展开到 ``style.font.size``。
        Word 中常见的“正文派生样式/标题派生样式”因此必须显式沿基类链读取；
        循环保护用于拒绝损坏或第三方生成器产生的异常样式关系。
        """

        sizes: list[Any] = []
        seen: set[object] = set()
        current = style
        while current is not None:
            style_element = getattr(current, "_element", None)
            style_id = getattr(current, "style_id", None)
            key: object = (
                ("style-id", style_id)
                if style_id is not None
                else ("element-id", id(style_element or current))
            )
            if key in seen:
                break
            seen.add(key)
            font = getattr(current, "font", None)
            sizes.append(getattr(font, "size", None) if font is not None else None)
            current = getattr(current, "base_style", None)
        return sizes

    @staticmethod
    def _effective_font_color(run: Run, paragraph: Paragraph) -> str | None:
        for font in (
            run.font,
            getattr(getattr(run, "style", None), "font", None),
            getattr(getattr(paragraph, "style", None), "font", None),
        ):
            if font is None:
                continue
            color = getattr(font, "color", None)
            if color is None:
                continue
            try:
                rgb = color.rgb
            except (AttributeError, TypeError, ValueError):
                rgb = None
            if rgb is not None:
                return f"rgb:{rgb}"
            try:
                theme = color.theme_color
            except (AttributeError, TypeError, ValueError):
                theme = None
            if theme is not None:
                return f"theme:{theme}"
        return None

    @staticmethod
    def _heading_level(style_name: str | None) -> int | None:
        normalized = " ".join(str(style_name or "").split()).casefold()
        if normalized == "title":
            return 1
        match = re.fullmatch(r"heading\s*([1-6])", normalized)
        return int(match.group(1)) if match else None

    def _list_info(self, paragraph: Paragraph, doc) -> tuple[bool, int] | None:
        style_name = " ".join(
            str(paragraph.style.name if paragraph.style else "").split()
        ).casefold()
        match = re.fullmatch(r"list\s*(bullet|number)\s*([0-9]*)", style_name)
        if match:
            level = max(int(match.group(2) or "1") - 1, 0)
            return match.group(1) == "number", level

        num_pr = paragraph._p.pPr.find(qn("w:numPr")) if paragraph._p.pPr is not None else None
        if num_pr is None:
            return None
        num_id = num_pr.find(qn("w:numId"))
        ilvl = num_pr.find(qn("w:ilvl"))
        if num_id is None or ilvl is None:
            return None
        try:
            num_id_value = int(num_id.get(qn("w:val")))
            level = int(ilvl.get(qn("w:val")))
        except (TypeError, ValueError):
            return None
        ordered = self._numbering_is_ordered(doc, num_id_value, level)
        return ordered, max(level, 0)

    @staticmethod
    def _numbering_is_ordered(doc, num_id: int, level: int) -> bool:
        numbering_part = getattr(doc.part, "numbering_part", None)
        if numbering_part is None:
            return True
        root = numbering_part.element
        num_element = next(
            (item for item in root.findall(qn("w:num")) if item.get(qn("w:numId")) == str(num_id)),
            None,
        )
        if num_element is None:
            return True
        abstract_id = num_element.find(qn("w:abstractNumId"))
        if abstract_id is None:
            return True
        abstract = next(
            (
                item
                for item in root.findall(qn("w:abstractNum"))
                if item.get(qn("w:abstractNumId")) == abstract_id.get(qn("w:val"))
            ),
            None,
        )
        if abstract is None:
            return True
        level_element = next(
            (
                item
                for item in abstract.findall(qn("w:lvl"))
                if item.get(qn("w:ilvl")) == str(level)
            ),
            None,
        )
        num_format = level_element.find(qn("w:numFmt")) if level_element is not None else None
        return num_format is None or num_format.get(qn("w:val")) != "bullet"

    def _rich_table(
        self,
        table_element,
        doc,
        image_parts: dict[str, Any],
        image_dir: Path,
        image_counter: list[int],
        image_paths: list[Path],
        block_counter: list[int],
    ) -> dict[str, Any]:
        rows: list[dict[str, Any]] = []
        for row_element in table_element.findall(qn("w:tr")):
            cells: list[dict[str, Any]] = []
            for cell_element in row_element.findall(qn("w:tc")):
                cell_blocks: list[dict[str, Any]] = []
                for child in cell_element:
                    if child.tag == qn("w:p"):
                        cell_blocks.append(
                            self._rich_paragraph(
                                Paragraph(child, doc),
                                doc,
                                image_parts,
                                image_dir,
                                image_counter,
                                image_paths,
                                block_counter,
                            )
                        )
                    elif child.tag == qn("w:tbl"):
                        cell_blocks.append(
                            self._rich_table(
                                child,
                                doc,
                                image_parts,
                                image_dir,
                                image_counter,
                                image_paths,
                                block_counter,
                            )
                        )
                cell: dict[str, Any] = {"blocks": cell_blocks}
                cell_properties = cell_element.find(qn("w:tcPr"))
                grid_span = (
                    cell_properties.find(qn("w:gridSpan")) if cell_properties is not None else None
                )
                if grid_span is not None:
                    try:
                        span = int(grid_span.get(qn("w:val")))
                    except (TypeError, ValueError):
                        span = 1
                    if span > 1:
                        cell["colspan"] = span
                cells.append(cell)
            rows.append({"cells": cells})
        return {
            "kind": "table",
            "block_id": self._next_block_id(block_counter),
            "rows": rows,
        }

    @staticmethod
    def _hyperlink_target(doc, hyperlink) -> tuple[str, str | None]:
        rel_id = hyperlink.get(qn("r:id"))
        relationship = doc.part.rels.get(rel_id) if rel_id else None
        target = getattr(relationship, "target_ref", None)
        if not target or not DocxParser._is_safe_http_url(target):
            raise ValueError("DOCX 超链接不是安全的 HTTP(S) 地址")
        tooltip = hyperlink.get(qn("w:tooltip"))
        return target, tooltip.strip() if tooltip and tooltip.strip() else None

    @staticmethod
    def _is_safe_http_url(value: str) -> bool:
        if any(ord(char) < 0x20 for char in value) or "\\" in value:
            return False
        parsed = urlsplit(value)
        return parsed.scheme in {"http", "https"} and bool(parsed.netloc)

    def _title_from_rich_blocks(self, blocks: list[dict[str, Any]], filename: str) -> str:
        for block in blocks:
            if block.get("kind") == "heading" and block.get("level") == 1:
                text = self._rich_block_text(block)
                if text.strip():
                    return self._normalize_title(text)
        for block in blocks:
            text = self._rich_block_text(block)
            if text.strip():
                return self._normalize_title(text)
        return Path(filename).stem

    def _leading_title_block(
        self,
        blocks: list[dict[str, Any]],
    ) -> dict[str, Any] | None:
        """返回开头三个非空文本块中可唯一绑定的正文主标题。

        当模板错误地把后续小节也标为 H1 时，第一非空块本身是 H1 仍是稳定
        的主标题证据；否则继续要求候选唯一，避免猜测位于正文中的 H1。
        """

        leading_text_blocks: list[dict[str, Any]] = []
        for block in blocks:
            if block.get("kind") not in {"paragraph", "heading"}:
                continue
            if not self._rich_block_text(block).strip():
                continue
            leading_text_blocks.append(block)
            if len(leading_text_blocks) == 3:
                break
        candidates = [
            block
            for block in leading_text_blocks
            if block.get("kind") == "heading" and block.get("level") == 1
        ]
        if candidates and leading_text_blocks[0] is candidates[0]:
            return candidates[0]
        return candidates[0] if len(candidates) == 1 else None

    @classmethod
    def _normalize_body_heading_one(
        cls,
        blocks: list[dict[str, Any]],
        title_block_id: str | None,
    ) -> None:
        """保证 canonical 文档只有已绑定标题可以保留 H1。"""

        for block in blocks:
            kind = block.get("kind")
            if (
                kind == "heading"
                and block.get("level") == 1
                and block.get("block_id") != title_block_id
            ):
                block["level"] = 2
                style_name = " ".join(str(block.get("style_name") or "").split())
                if style_name.casefold() in {"heading 1", "title"}:
                    block["style_name"] = "Heading 2"
            elif kind == "list":
                for item in block.get("items", []):
                    cls._normalize_body_heading_one(
                        item.get("blocks", []),
                        title_block_id,
                    )
            elif kind == "table":
                for row in block.get("rows", []):
                    for cell in row.get("cells", []):
                        cls._normalize_body_heading_one(
                            cell.get("blocks", []),
                            title_block_id,
                        )

    def _find_title_block_id(self, blocks: list[dict[str, Any]], title: str) -> str | None:
        normalized_title = self._normalize_title(title)
        if not normalized_title:
            return None
        for block in blocks:
            if block.get("kind") not in {"paragraph", "heading"}:
                continue
            if self._normalize_title(self._rich_block_text(block)) == normalized_title:
                return block.get("block_id")
        return None

    @staticmethod
    def _rich_block_text(block: dict[str, Any]) -> str:
        if block.get("kind") in {"paragraph", "heading"}:
            return "".join(
                child.get("text", "")
                for child in block.get("children", [])
                if child.get("kind") == "text"
            )
        if block.get("kind") == "list":
            return "".join(
                DocxParser._rich_block_text(nested)
                for item in block.get("items", [])
                for nested in item.get("blocks", [])
            )
        if block.get("kind") == "table":
            return "".join(
                DocxParser._rich_block_text(nested)
                for row in block.get("rows", [])
                for cell in row.get("cells", [])
                for nested in cell.get("blocks", [])
            )
        return ""

    def _legacy_blocks_from_rich(self, blocks: list[dict[str, Any]]) -> list[ContentBlock]:
        result: list[ContentBlock] = []

        def visit(block: dict[str, Any]) -> None:
            kind = block.get("kind")
            if kind in {"paragraph", "heading"}:
                for child in block.get("children", []):
                    if child.get("kind") == "text":
                        result.append(
                            ContentBlock(
                                type="heading" if kind == "heading" else "text",
                                text=child.get("text"),
                                style_name=block.get("style_name"),
                                position=len(result),
                            )
                        )
                    elif child.get("kind") == "image":
                        path = child.get("__source_path")
                        result.append(
                            ContentBlock(
                                type="image",
                                image_filename=child.get("__original_filename"),
                                image_path=path,
                                image_width=child.get("width"),
                                image_height=child.get("height"),
                                position=len(result),
                            )
                        )
            elif kind == "list":
                for item in block.get("items", []):
                    for nested in item.get("blocks", []):
                        visit(nested)
            elif kind == "table":
                for row in block.get("rows", []):
                    for cell in row.get("cells", []):
                        for nested in cell.get("blocks", []):
                            visit(nested)

        for block in blocks:
            visit(block)
        return result

    def _find_paragraph(self, doc, element):
        """根据 XML 元素查找对应的 python-docx Paragraph 对象"""
        for para in doc.paragraphs:
            if para._element is element:
                return para
        return None

    def _paragraph_parts(self, para, image_parts):
        """Return text/image parts while preserving paragraph XML order."""
        parts = []
        text_buffer = []

        def flush_text() -> None:
            if not text_buffer:
                return
            text = "".join(text_buffer)
            text_buffer.clear()
            if text.strip():
                parts.append(("text", text))

        for run in para.runs:
            for node in run._element.iter():
                if node.tag == W_TEXT:
                    text_buffer.append(node.text or "")
                elif node.tag == W_TAB:
                    text_buffer.append("\t")
                elif node.tag in {W_BREAK, W_CARRIAGE_RETURN}:
                    text_buffer.append("\n")
                elif node.tag == A_BLIP:
                    embed = node.get(f"{{{R_NS}}}embed")
                    image_part = image_parts.get(embed)
                    if image_part is not None:
                        flush_text()
                        parts.append(("image", image_part))

        flush_text()
        return parts

    @staticmethod
    def _normalize_title(value: str | None) -> str:
        """Normalize title whitespace without changing non-whitespace text."""
        return " ".join(str(value or "").split())

    @staticmethod
    def _is_heading_style(style_name: str | None) -> bool:
        normalized = " ".join(str(style_name or "").split()).casefold()
        return normalized == "title" or bool(re.fullmatch(r"heading\s*[1-9]", normalized))

    @staticmethod
    def _is_document_title_style(style_name: str | None) -> bool:
        normalized = " ".join(str(style_name or "").split()).casefold()
        return normalized == "title" or bool(re.fullmatch(r"heading\s*1", normalized))

    def _extract_table_text(self, tbl_element) -> str:
        """提取表格文本"""
        rows = tbl_element.findall(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr"
        )
        lines = []
        for row in rows:
            cells = row.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc")
            cell_texts = []
            for cell in cells:
                paras = cell.findall(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
                )
                cell_text = " ".join(
                    "".join(
                        t.text or ""
                        for t in p.iterfind(
                            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
                        )
                    )
                    for p in paras
                )
                cell_texts.append(cell_text.strip())
            lines.append(" | ".join(cell_texts))
        return "\n".join(lines)

    def _get_image_size(self, img_path: str) -> tuple:
        """获取图片尺寸"""
        try:
            with Image.open(img_path) as img:
                return img.size
        except Exception:
            return (0, 0)

    def to_json(self, article: ParsedArticle) -> str:
        """将解析结果转为 JSON"""
        data = {
            "filename": article.filename,
            "original_path": article.original_path,
            "plain_text": article.plain_text,
            "char_count": article.char_count,
            "image_count": article.image_count,
            "blocks": [
                {
                    "type": b.type,
                    "text": b.text,
                    "style_name": b.style_name,
                    "image_filename": b.image_filename,
                    "image_path": b.image_path,
                    "image_width": b.image_width,
                    "image_height": b.image_height,
                    "position": b.position,
                }
                for b in article.blocks
            ],
        }
        return json.dumps(data, ensure_ascii=False, indent=2)

    def get_title_from_article(self, article: ParsedArticle) -> str:
        """Extract the selected title using core properties/style/body priority."""
        core_title = article.document_title or ""
        if core_title:
            return core_title

        for block in article.blocks:
            if (
                block.type == "heading"
                and self._is_document_title_style(block.style_name)
                and block.text
            ):
                return self._normalize_title(block.text)

        # Fall back to the first non-empty body block, without truncating it.
        for block in article.blocks:
            if block.type == "text" and block.text and block.text.strip():
                return self._normalize_title(block.text)

        return os.path.splitext(article.filename)[0]
